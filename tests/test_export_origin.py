"""Every exported row states its Origin.

Carved rows always said "Carved (deleted/unflushed)", but build_entry_row
never set Origin for a LIVE row, so its cell was blank in the CSV, the xlsx
"All Records" sheet and the sqlite `records` table -- while the screen said
"Live record". A consumer filtering on Origin could not tell live from
unknown: a blank standing in for a known value.
"""
import csv
import sqlite3

import openpyxl

from webcache import provenance
from webcache.engine import forensic_engine


def _live():
    # A raw parser entry: no Origin key of its own, exactly as the engine
    # holds it.
    return [{"Url": "http://example.test/a", "ContainerType": "History",
             "ContainerName": "History", "EntryId": 1, "ContainerId": 3}]


def _recovered():
    return [
        {"Url": "http://carved.test/1", "RecoverySource": "log",
         "Confidence": "medium"},
        {"Url": "", "RecoverySource":
         "Deleted-record stub (b-tree key only)", "Confidence": "exact key",
         "TableName": "Container_3", "EntryId": 7},
    ]


def _engine():
    engine = forensic_engine.ForensicEngine()
    engine.load_entries(_live())
    return engine


def test_live_record_origin_constant_exists():
    assert getattr(provenance, "LIVE_ORIGIN", None) == "Live record"


def test_csv_live_rows_say_live_and_carved_rows_say_carved(tmp_path):
    path = str(tmp_path / "a.csv")
    _engine().export_csv(_live(), path, recovered=_recovered())
    with open(path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    col = rows[0].index("Origin")
    origins = [r[col] for r in rows[1:]]
    assert origins == ["Live record", provenance.CARVED_ORIGIN,
                       provenance.CARVED_ORIGIN], origins


def test_xlsx_all_records_live_row_says_live(tmp_path):
    engine = _engine()
    path = str(tmp_path / "b.xlsx")
    engine.export_xlsx(_live(), path, True, {"examiner": "E"}, {}, {},
                       _recovered(), None, None, engine.get_statistics())
    ws = openpyxl.load_workbook(path)["All Records"]
    heads = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=heads.index("Origin") + 1).value == \
        "Live record"


def test_sqlite_records_origin_is_never_blank(tmp_path):
    path = str(tmp_path / "c.sqlite")
    _engine().export_sqlite(_live(), path, recovered=_recovered())
    con = sqlite3.connect(path)
    try:
        got = [r[0] for r in con.execute(
            "SELECT Origin FROM records ORDER BY id")]
    finally:
        con.close()
    assert got == ["Live record", provenance.CARVED_ORIGIN,
                   provenance.CARVED_ORIGIN], got
