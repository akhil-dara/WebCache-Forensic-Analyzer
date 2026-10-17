"""Hostile evidence through every surface of the running page (spec:
"Security model" 3-5). The file half is tests/test_html_report_escaping.py.

Every value below is synthetic. A probe, installed before the page's own
script runs (Page.addScriptToEvaluateOnNewDocument), records:
  * every element inserted whose tag or attribute the page never makes,
    and any inline handler or `data-wcpwn` marker at all;
  * every call to a markup-parsing API (innerHTML, insertAdjacentHTML,
    document.write, DOMParser, createContextualFragment ...), which the
    page's own code never makes;
  * every createElement of a tag the page never creates;
  * every Content-Security-Policy violation the page reports.
The live payloads set `window.__wc_pwned`; it must stay undefined.

A clean result means nothing unless the same instruments can see an
injection, so every surface has a POSITIVE CONTROL: once the surface has
been checked, a live payload is planted into that very surface through
innerHTML and every detector must fire. The downloads' checkers, the
island/title scan, the canary and the network log each have their own
planted control in the last test."""
import csv
import json
import os
import pathlib
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

import pytest

from webcache.engine import forensic_engine
from webcache.shell import export_formats
from webcache.shell.html_package import (PackageColumn, PackageView,
                                         build_island, encode_cell,
                                         island_json, render_page)
from webcache.shell.record_kinds import K_PERIOD
from webcache.views.records_view import member_record_columns

from tests import _package_fixtures as fx
from tests.test_html_report_escaping import HOSTILE

CANARY = "window.__wc_pwned=1"
#: One megabyte, in 10,000 lines, led by a marker element.
MEGA = "<wcx-pwn data-wcpwn>" + ("x" * 99 + "\n") * 10000
PAYLOADS = HOSTILE + [
    "<script>%s</script>" % CANARY,
    "<img src=x data-wcpwn onerror=\"%s\">" % CANARY,
    "<svg data-wcpwn onload=\"%s\"></svg>" % CANARY,
    "\" data-wcpwn onmouseover=\"%s\" x=\"" % CANARY,
    "' data-wcpwn autofocus onfocus='%s' x='" % CANARY,
    "\"><wcx-pwn data-wcpwn></wcx-pwn>",
    "javascript:%s" % CANARY,
    "data:text/html,<script>%s</script>" % CANARY,
    "&lt;img src=x data-wcpwn onerror=%s&gt;" % CANARY,
    "&#60;svg data-wcpwn onload=%s&#62;" % CANARY,
    "</script><script>%s</script>" % CANARY,
    "<!--<script>%s</script>" % CANARY,
    "]]><wcx-pwn data-wcpwn/><![CDATA[",
    "</title><wcx-pwn data-wcpwn></wcx-pwn>",
    "</style><wcx-pwn data-wcpwn></wcx-pwn>",
    "</textarea><wcx-pwn data-wcpwn></wcx-pwn>",
    "@@SCRIPT@@ @@DATA@@ @@CSP@@ @@STYLE@@",
    "a\u2028b\u2029c",
    "\u202egnp.exe\u2066x\u2069\u200f",
    "a\x00b",
    "a\ud800b",
    "a\udfffb",
    "a\ufffeb\uffffc",
    "=HYPERLINK(\"https://example.test\",\"x\")",
    "=cmd|' /C calc'!A0",
    "+1+1",
    "-1+1",
    "@SUM(1,1)",
    "\t=1+1",
    "\r=1+1",
    "\n".join(["<img src=x data-wcpwn onerror=%s>" % CANARY] * 10),
    MEGA,
]
#: What the island stores for each: forensic_engine._safe_text's visible
#: escapes for NUL and lone surrogates; everything else exactly as given.
STORED = [encode_cell(p) for p in PAYLOADS]
SHORT = [s for s in STORED if len(s) < 200]
#: What a time column stores for the text placed in it (_hostile_snapshot).
TIME_TEXT = [encode_cell(PAYLOADS[i % len(PAYLOADS)], True)
             for i in range(3, len(PAYLOADS) + 1, 3)]


def as_label(text):
    """What the page shows for a string that is not a cell (a label, a
    summary value, a path): _safe_text's output, as every other writer
    writes it -- a control character, lone surrogate, DEL or noncharacter
    as its visible escape."""
    return forensic_engine._safe_text(text)

#: A live payload, planted into a surface by the POSITIVE CONTROLS.
PLANT = "<img src=x data-wcpwn onerror=\"%s\">" % CANARY

#: Every tag the page's code or template puts in the document, and every
#: attribute it sets. Anything else, any on* attribute and the
#: data-wcpwn marker are an injection.
OK_TAGS = ("html head meta title style body noscript script div span button "
           "table colgroup col thead tbody tr td th input select option "
           "optgroup label h1 h2 h3 p pre b strong i mark ol li header nav "
           "main textarea details summary").split()
#: createElement: the page also makes a detached <a> (saveBlob) and a
#: detached <canvas> (text measuring); neither is ever inserted.
OK_CREATED = OK_TAGS + ["a", "canvas"]
OK_ATTRS = ("lang charset name content http-equiv type id class title role "
            "style tabindex placeholder hidden disabled scope colspan value "
            # data-col: a column's numeric index on its heading buttons
            # (grid.js), never evidence text.
            # open: the top bar's evidence-hash list when shown.
            "label data-chip data-pos data-theme data-col open").split()
#: Elements only the template may hold: one more of any is an injection.
TEMPLATE_CENSUS = ("script,style,meta,title,link,base,iframe,object,embed,"
                   "img,svg,math,form,a,frame,video,audio,source,wcx-pwn")

PROBE = r"""(function () {
  var OK_TAGS = new Set(%(tags)s), OK_CREATED = new Set(%(created)s);
  var OK_ATTRS = new Set(%(attrs)s);
  var LATE = new Set(["script", "style", "meta", "title", "link", "base"]);
  var probe = window.__wc_probe = {flags: [], parses: [], created: [], csp: []};
  function attrBad(name) {
    name = name.toLowerCase();
    return name.indexOf("on") === 0 || name === "data-wcpwn" ||
      !(OK_ATTRS.has(name) || name.indexOf("aria-") === 0);
  }
  function check(node) {
    if (!node || node.nodeType !== 1) return;
    var tag = node.nodeName.toLowerCase();
    if (!OK_TAGS.has(tag)) probe.flags.push("tag " + tag);
    if (LATE.has(tag) && document.readyState !== "loading") probe.flags.push("late " + tag);
    for (var i = 0; i < node.attributes.length; i++) {
      if (attrBad(node.attributes[i].name)) probe.flags.push("attr " + node.attributes[i].name + " on " + tag);
    }
    for (var c = node.firstElementChild; c; c = c.nextElementSibling) check(c);
  }
  new MutationObserver(function (records) {
    records.forEach(function (r) {
      if (r.type === "childList") r.addedNodes.forEach(check);
      else if (r.type === "attributes" && attrBad(r.attributeName)) {
        probe.flags.push("attr " + r.attributeName + " set on " + r.target.nodeName.toLowerCase());
      }
    });
  }).observe(document, {childList: true, subtree: true, attributes: true});
  document.addEventListener("securitypolicyviolation", function (e) {
    probe.csp.push(e.violatedDirective + " " + e.blockedURI);
  }, true);
  function wrapSetter(proto, name) {
    var d = proto && Object.getOwnPropertyDescriptor(proto, name);
    if (!d || !d.set) return;
    Object.defineProperty(proto, name, {configurable: true, enumerable: d.enumerable, get: d.get,
      set: function (v) { probe.parses.push(name); return d.set.call(this, v); }});
  }
  function wrapMethod(proto, name) {
    var f = proto && proto[name];
    if (typeof f !== "function") return;
    proto[name] = function () { probe.parses.push(name); return f.apply(this, arguments); };
  }
  wrapSetter(Element.prototype, "innerHTML");
  wrapSetter(Element.prototype, "outerHTML");
  wrapSetter(window.ShadowRoot && ShadowRoot.prototype, "innerHTML");
  wrapMethod(Element.prototype, "insertAdjacentHTML");
  wrapMethod(Element.prototype, "setHTMLUnsafe");
  wrapMethod(Document.prototype, "write");
  wrapMethod(Document.prototype, "writeln");
  wrapMethod(Range.prototype, "createContextualFragment");
  wrapMethod(DOMParser.prototype, "parseFromString");
  var create = Document.prototype.createElement;
  Document.prototype.createElement = function (tag) {
    if (!OK_CREATED.has(String(tag).toLowerCase())) probe.created.push(String(tag));
    return create.apply(this, arguments);
  };
})();""" % {"tags": json.dumps(OK_TAGS), "created": json.dumps(OK_CREATED),
            "attrs": json.dumps(OK_ATTRS)}

PROBE_STATE = "JSON.stringify(window.__wc_probe)"
PROBE_RESET = ("(function (p) { p.flags.length = p.parses.length = "
               "p.created.length = p.csp.length = 0; })(window.__wc_probe)")


# --- the hostile snapshot -------------------------------------------------------

def _hostile_snapshot():
    rows = []
    for i, payload in enumerate(PAYLOADS, 1):
        carved = bool(i % 2)
        row = fx.history_row(i, carved=carved, CleanUrl=payload)
        record = row["_record"]
        record["Url"] = record["CleanUrl"] = record["Domain"] = payload
        if carved:
            record["Confidence"] = row["Source"] = PAYLOADS[(i + 3) % len(PAYLOADS)]
        if i % 3 == 0:
            # A time column holding text: the date tree's "other" leaves.
            row["AccessedUtc"] = PAYLOADS[i % len(PAYLOADS)]
        rows.append(row)
    cols = fx.history_columns()
    cols[0].label, cols[0].note = PAYLOADS[8], PAYLOADS[9]
    cols[0].badge, cols[0].badge_tooltip = "OBS", PAYLOADS[10]
    quick = [{"label": PAYLOADS[11], "key": "IsInPrivate", "values": ["Yes"],
              "cond": None, "sentence": PAYLOADS[12]}]
    seg = fx.segment("history", PAYLOADS[13], rows=rows, columns=cols,
                     quick=quick, note=PAYLOADS[14], caveat=PAYLOADS[15],
                     subtitle=PAYLOADS[16])
    members = [fx.entry(100 + k, Url=p, CleanUrl=p)
               for k, p in enumerate(PAYLOADS[:12])]
    period = {"Period": PAYLOADS[17], "Records": len(members),
              "_kind": K_PERIOD, "_members": members}
    day = fx.segment("day_utc", PAYLOADS[18], rows=[period],
                     columns=[PackageColumn("Period", PAYLOADS[19]),
                              PackageColumn("Records", "Records",
                                            kind="num")])
    view = PackageView("history", PAYLOADS[20], PAYLOADS[21], PAYLOADS[22],
                       [seg, day])
    snap = fx.snapshot(
        views=[view],
        summary=[(PAYLOADS[23], PAYLOADS[24]), ("File", PAYLOADS[25])],
        scope=[(PAYLOADS[26], PAYLOADS[27])],
        basis=[("OBS", PAYLOADS[28], PAYLOADS[29])],
        member_columns=[PackageColumn.from_column(c)
                        for c in member_record_columns()])
    snap.meta["evidence"] = [{"path": "C:\\case\\" + PAYLOADS[9],
                              "sha256": PAYLOADS[7], "bytes": 1}]
    snap.meta["case"] = {"examiner": PAYLOADS[30], PAYLOADS[8]: "alice"}
    snap.integrity = {"hashes": [[PAYLOADS[31], PAYLOADS[32], PAYLOADS[33]]],
                      "source": [[PAYLOADS[34], PAYLOADS[35], ""]],
                      "log": [[PAYLOADS[36], True], [PAYLOADS[1], False]]}
    return snap


# --- the instruments -------------------------------------------------------------

def _census(browser):
    return browser.evaluate(
        "(function () { var out = {}; %s.split(',').forEach(function (t) {"
        " out[t] = document.getElementsByTagName(t).length; }); return out; })()"
        % json.dumps(TEMPLATE_CENSUS))


def _settle(browser, ms=40):
    browser.evaluate("new Promise(function (ok) { requestAnimationFrame(function ()"
                     " { setTimeout(ok, %d); }); })" % ms, await_promise=True)


def _probe(browser):
    return json.loads(browser.evaluate(PROBE_STATE))


def _csp_since(browser, mark):
    return [e["params"]["entry"].get("text", "")
            for e in browser.events[mark:]
            if e.get("method") == "Log.entryAdded"
            and "Content Security Policy" in e["params"]["entry"].get("text", "")]


def _requests_since(browser, mark, page_url):
    """Requests since `mark` other than the page itself and the blob:
    URLs of its own downloads."""
    return [e["params"]["request"]["url"] for e in browser.events[mark:]
            if e.get("method") == "Network.requestWillBeSent"
            and e["params"]["request"]["url"] != page_url
            and not e["params"]["request"]["url"].startswith("blob:")]


def _errors_since(browser, mark):
    keep = browser.events
    try:
        browser.events = keep[mark:]
        return browser.console_errors()
    finally:
        browser.events = keep


def _js(value):
    return json.dumps(value)


class Surfaces:
    """Walks the hostile page one surface at a time."""

    def __init__(self, browser, clean_census, page_url):
        self.browser = browser
        self.clean = clean_census
        self.page_url = page_url
        self.checked = []

    def check(self, name, action, shows, plant_into):
        """`action` puts hostile evidence on screen; `shows` is a JS
        expression that is true only when that evidence is on screen AS
        TEXT (so a surface that drew nothing cannot pass); `plant_into` is
        the selector of the surface's own node, for the positive control."""
        b = self.browser
        mark = len(b.events)
        if action:
            b.evaluate(action, await_promise=True)
        _settle(b)
        assert b.evaluate(shows) is True, "%s: the hostile text is not shown" % name
        state = _probe(b)
        assert state == {"flags": [], "parses": [], "created": [], "csp": []}, (name, state)
        assert _census(b) == self.clean, name
        assert b.evaluate("typeof window.__wc_pwned") == "undefined", name
        assert _csp_since(b, mark) == [], name
        assert _errors_since(b, mark) == [], name
        assert _requests_since(b, mark, self.page_url) == [], name
        # Positive control: the same instruments, a live payload, this node.
        mark = len(b.events)
        b.evaluate("(function () { var host = document.querySelector(%s);"
                   " var s = document.createElement('span'); s.id = 'wc-plant';"
                   " host.appendChild(s); s.innerHTML = %s; })()"
                   % (_js(plant_into), _js(PLANT)))
        _settle(b, 120)
        state = _probe(b)
        assert any("data-wcpwn" in f for f in state["flags"]), (name, state)
        assert any(f.startswith("tag img") for f in state["flags"]), (name, state)
        assert "innerHTML" in state["parses"], (name, state)
        assert state["csp"], (name, state)
        assert _csp_since(b, mark), name
        assert _census(b)["img"] == self.clean["img"] + 1, name
        # The request log sees the planted image ask for its source (the
        # CSP then refuses it).
        assert any(u.endswith("/x") for u in
                   _requests_since(b, mark, self.page_url)), name
        b.evaluate("document.getElementById('wc-plant').remove()")
        b.evaluate(PROBE_RESET)
        self.checked.append(name)


def _shows_all(selector, texts):
    """JS: every one of `texts` is somewhere in the text of `selector`."""
    return ("(function () { var t = Array.from(document.querySelectorAll(%s))"
            ".map(function (n) { return n.textContent; }).join('\\u0000');"
            " return %s.every(function (x) { return t.indexOf(x) >= 0; }); })()"
            % (_js(selector), _js(texts)))


def _escape(browser):
    browser.evaluate("document.dispatchEvent(new KeyboardEvent('keydown',"
                     " {key: 'Escape'}))")


# --- the downloads' checkers -----------------------------------------------------

def _csv_rows(path):
    old = csv.field_size_limit()
    # The default reader refuses a field over 128 KiB; a 1 MB value is a
    # legitimate field, so the limit is raised for this read only.
    csv.field_size_limit(sys.maxsize if sys.maxsize < 2 ** 31 else 2 ** 31 - 1)
    try:
        with open(path, encoding="utf-8-sig", newline="") as fh:
            return list(csv.reader(fh, strict=True))
    finally:
        csv.field_size_limit(old)


def _csv_formula_cells(rows):
    """Cells Excel would evaluate on opening the CSV: the four leads, as
    export_formats._is_formula_text counts them."""
    return sum(1 for row in rows[1:] for v in row
               if export_formats._is_formula_text(v))


def _xlsx_formulas(path):
    """(formula cells openpyxl reads, <f> elements in the sheet XML). Every
    XML part must parse: a character XML 1.0 cannot carry fails here."""
    import openpyxl
    tags = 0
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        for name in z.namelist():
            data = z.read(name)
            assert b"<!DOCTYPE" not in data and b"<!ENTITY" not in data, name
            root = ET.fromstring(data)
            tags += sum(1 for e in root.iter() if e.tag.endswith("}f"))
    wb = openpyxl.load_workbook(path)
    cells = sum(1 for ws in wb.worksheets for row in ws.iter_rows()
                for c in row if c.data_type == "f")
    return cells, tags, wb


def _xlsx_text(stored):
    """What the page's workbook writes for a stored string: a character XML
    cannot carry as its visible escape, then Excel's cell limit stated."""
    text = stored.replace("\ufffe", "\\uFFFE").replace("\uffff", "\\uFFFF")
    return forensic_engine.xlsx_cell_parts(text, False)[0]


# --- the tests -------------------------------------------------------------------

@pytest.fixture()
def probed(browser):
    ident = browser.call("Page.addScriptToEvaluateOnNewDocument",
                         source=PROBE)["identifier"]
    try:
        yield browser
    finally:
        browser.call("Page.removeScriptToEvaluateOnNewDocument",
                     identifier=ident)
        try:
            browser.evaluate("WC.setWrap(true)")
        except RuntimeError:
            pass


def test_no_surface_turns_hostile_evidence_into_markup(probed, tmp_path):
    b = probed
    b.open(fx.write(tmp_path, None, name="clean.html"))
    clean = _census(b)
    clean_probe = _probe(b)
    # Negative control: a clean page trips none of the instruments.
    assert clean_probe == {"flags": [], "parses": [], "created": [], "csp": []}
    assert b.csp_violations() == []
    b.clear_downloads()
    snap = _hostile_snapshot()
    page = fx.write(tmp_path, snap, name="hostile.html")
    page_url = pathlib.Path(page).resolve().as_uri()
    b.open(page)
    s = Surfaces(b, clean, page_url)

    # The island: the parser ended it where the page did, and every value
    # in it reads back exactly as stored.
    assert b.evaluate("document.scripts.length") == 2
    assert b.evaluate("WC.data().views[0].segments[0].rows.map("
                      "function (r) { return r[0]; })") == STORED
    assert b.evaluate("document.querySelectorAll('[data-wcpwn], wcx-pwn').length") == 0

    # The title: evidence reaches it, as text.
    base = re.split(r"[\\/]", snap.meta["evidence"][0]["path"])[-1]
    assert b.evaluate("document.title") == "WebCache offline report \u2014 " + base
    assert b.evaluate("document.querySelector('title').children.length") == 0
    s.check("title and top bar", None,
            _shows_all("header.topbar", [base, PAYLOADS[7]]), "header.topbar")

    summary = [PAYLOADS[23], PAYLOADS[24], PAYLOADS[25], PAYLOADS[26],
               PAYLOADS[27], PAYLOADS[28], PAYLOADS[29], PAYLOADS[30],
               PAYLOADS[8], PAYLOADS[31], PAYLOADS[32], PAYLOADS[33],
               PAYLOADS[34], PAYLOADS[35], PAYLOADS[36], PAYLOADS[20],
               PAYLOADS[13], PAYLOADS[18], snap.meta["evidence"][0]["path"]]
    s.check("case summary", "WC.go('summary')",
            _shows_all("#wc-main", [as_label(x) for x in summary]),
            "#wc-main .page")
    s.check("navigation", None,
            _shows_all("nav.side", [PAYLOADS[20], PAYLOADS[21]]), "nav.side")

    s.check("view head: caveat, subtitle, count note, segment buttons",
            "WC.openSegment('history', 'history')",
            _shows_all(".vhead", [PAYLOADS[20], PAYLOADS[15], PAYLOADS[16],
                                  PAYLOADS[14], PAYLOADS[13], PAYLOADS[18]]),
            ".vhead")
    s.check("column heading, note and badge", None,
            "(function () { var th = document.querySelector('#wc-grid thead th:nth-child(2)');"
            " return th.textContent.indexOf(%s) >= 0 && th.title === %s &&"
            " th.querySelector('.badge').title === %s; })()"
            % (_js(PAYLOADS[8]), _js(PAYLOADS[9]), _js(PAYLOADS[10])),
            "#wc-grid thead th")

    # Grid cells: scroll the whole virtualised grid, every row drawn once.
    seen = set()
    for _ in range(400):
        seen.update(b.evaluate(
            "Array.from(document.querySelectorAll('#wc-grid tbody tr:not(.spacer)'))"
            ".map(function (tr) { return tr.cells[1].firstChild.textContent; })"))
        if seen >= set(STORED):
            break
        b.evaluate("(function (w) { w.scrollTop += Math.max(200, w.clientHeight - 40); })"
                   "(document.getElementById('wc-grid'))")
        _settle(b, 10)
    assert seen == set(STORED)
    b.evaluate("document.getElementById('wc-grid').scrollTop = 0")
    s.check("grid cells", None, _shows_all("#wc-grid tbody", SHORT[:6]),
            "#wc-grid tbody td .cell")

    s.check("quick chip and the sentence it writes",
            "document.querySelector('button.chip').click()",
            "document.querySelector('button.chip').textContent === %s && %s"
            % (_js(PAYLOADS[11]), _shows_all("#wc-sentence", [PAYLOADS[12]])),
            "#wc-sentence")
    b.evaluate("document.querySelector('button.chip').click()")

    s.check("filter popup: value list",
            "document.querySelectorAll('#wc-grid th .funnel')[0].click()",
            "(function () { var v = Array.from(document.querySelectorAll('.popup .values label'));"
            " return v.length > 5 && v.every(function (l) { var t = l.querySelector('.v').textContent;"
            " return l.title === t && %s.indexOf(t) >= 0; }) &&"
            " document.querySelector('.popup strong').textContent === %s; })()"
            % (_js(STORED), _js(PAYLOADS[8])),
            ".popup")
    # Tick only the first value, apply: the sentence names it.
    first = b.evaluate("document.querySelector('.popup .values label .v').textContent")
    b.evaluate("(function () { var p = document.querySelector('.popup');"
               " Array.from(p.querySelectorAll('button')).filter(function (x) {"
               " return x.textContent === 'Select none'; })[0].click();"
               " p.querySelector('.values input').click();"
               " Array.from(p.querySelectorAll('button')).filter(function (x) {"
               " return x.textContent === 'Apply'; })[0].click(); })()")
    s.check("filter chips: a value filter's sentence", None,
            _shows_all("#wc-sentence", [first]), "#wc-sentence")
    b.evaluate("Array.from(document.querySelectorAll('.vhead button')).filter("
               "function (x) { return x.textContent === 'Clear filters'; })[0].click()")

    s.check("filter popup: date tree",
            "document.querySelectorAll('#wc-grid th .funnel')[1].click()",
            "(function () { var t = Array.from(document.querySelectorAll("
            "'.popup .values label .v')).map(function (n) { return n.textContent; });"
            " var odd = t.filter(function (x) { return !/^\\d{4}/.test(x); });"
            " return odd.length >= 3 && odd.every(function (x) { return %s.indexOf(x) >= 0; }); })()"
            % _js(TIME_TEXT), ".popup")
    _escape(b)

    s.check("filter box and its sentence",
            "WC.openSegment('history', 'history', %s)" % _js(PAYLOADS[8]),
            "document.getElementById('wc-filter').value === %s && %s"
            % (_js(PAYLOADS[8].strip()), _shows_all("#wc-sentence", [PAYLOADS[8].strip()])),
            "#wc-sentence")

    s.check("Search all views: input, snippets, segment buttons",
            "WC.go('search', '<img')",
            "document.getElementById('wc-search').value === '<img' && %s"
            % _shows_all(".results", ["<img", PAYLOADS[20], PAYLOADS[13]]),
            ".results .hit")
    s.check("Search all views: a hostile needle with no match",
            "WC.go('search', %s)" % _js(PAYLOADS[1] + "zz-no-such"),
            _shows_all("#wc-main .count", [PAYLOADS[1] + "zz-no-such"]),
            "#wc-main .page")
    s.check("Search all views: a hostile needle that matches",
            "WC.go('search', %s)" % _js(PAYLOADS[15]),
            _shows_all(".results .snip mark", [PAYLOADS[15]]), ".results")

    # Wrapped values: every "+N more" marker opens the whole value.
    b.evaluate("WC.openSegment('history', 'history', 'data-wcpwn')")
    b.evaluate("WC.setWrap(true)")
    _settle(b)
    markers = b.evaluate("document.querySelectorAll('#wc-grid .more').length")
    assert markers >= 2
    for k in range(markers):
        s.check("whole-value dialog %d" % k,
                "document.querySelectorAll('#wc-grid .more')[%d].click()" % k,
                "(function () { var p = document.querySelector('.dialog pre.whole');"
                " return !!p && %s.indexOf(p.textContent) >= 0 &&"
                " document.querySelector('.dialog h2').textContent === %s; })()"
                % (_js(STORED), _js(PAYLOADS[8])),
                ".dialog .dbody")
        _escape(b)

    # The record dialog, for every row: identity, strip, fields.
    heads = b.evaluate("(function () { var D = WC.data(), s = D.views[0].segments[0];"
                       " return s.hd.map(function (h) { return D.heads[h]; }); })()")
    for r, stored in enumerate(STORED):
        strip = [f[0] + ": " + f[1] for f in heads[r]["p"]]
        shows = ("(function () { var d = document.querySelector('.dialog');"
                 " if (!d || d.querySelector('h2').textContent !== %s) return false;"
                 " var t = d.textContent; if (%s.some(function (x) { return t.indexOf(x) < 0; }))"
                 " return false; var s = d.querySelector('.strip');"
                 " return %s ? !!s && %s.every(function (x) { return s.textContent.indexOf(x) >= 0; })"
                 " : !s; })()"
                 % (_js(heads[r]["i"]), _js([stored[:4000]]), _js(bool(strip)),
                    _js(strip)))
        s.check("record dialog, row %d" % r,
                "WC.openRecord('history', 'history', %d)" % r, shows,
                ".dialog .strip" if strip else ".dialog .kv .v")
        _escape(b)

    s.check("aggregate: the members table",
            "WC.openRecord('history', 'day_utc', 0)",
            "document.querySelector('.dialog h2').textContent === %s && %s"
            % (_js(PAYLOADS[19] + ": " + PAYLOADS[17]),
               _shows_all(".dialog table.members",
                                             [encode_cell(p) for p in PAYLOADS[:12]
                                              if p.strip()][:6])),
            ".dialog table.members td")
    s.check("aggregate: a member's own dialog",
            "document.querySelector('.dialog table.members button').click()",
            "document.querySelectorAll('.dialog').length === 2 &&"
            " document.querySelectorAll('.dialog')[1].textContent.indexOf(%s) >= 0"
            % _js(STORED[0]),
            ".dialog:last-of-type .dbody")
    _escape(b)
    _escape(b)

    # CSV, its provenance note, and the workbook: the whole segment.
    b.evaluate("WC.openSegment('history', 'history', '')")
    s.check("CSV export and the provenance note offered for it",
            "WC.exportCsv('history', 'history')",
            "document.getElementById('wc-toast').textContent.indexOf('Wrote ') >= 0",
            "#wc-toast")
    csv_path = b.wait_download("webcache_history_history.csv")
    # Saving the note dismisses the toast: the note's own
    # content is checked as a file below.
    b.evaluate("Array.from(document.querySelectorAll('#wc-toast button')).filter("
               "function (x) { return x.textContent === 'Save provenance note'; })[0].click()")
    note_path = b.wait_download("webcache_history_history.csv.provenance.txt")
    s.check("XLSX export", "WC.exportXlsx('history', 'history')",
            "document.getElementById('wc-toast').textContent.indexOf('.xlsx') >= 0",
            "#wc-toast")
    xlsx_path = b.wait_download("webcache_history_history.xlsx")

    # A filtered export: the hostile filter text goes into the note.
    mark = len(b.events)
    b.evaluate("WC.openSegment('history', 'history', %s)" % _js(PAYLOADS[15]))
    b.evaluate("WC.exportCsv('history', 'history')", await_promise=True)
    b.wait_download("webcache_history_history_filtered.csv")
    b.evaluate("Array.from(document.querySelectorAll('#wc-toast button')).filter("
               "function (x) { return x.textContent === 'Save provenance note'; })[0].click()")
    filtered_note = b.wait_download(
        "webcache_history_history_filtered.csv.provenance.txt")

    # Nothing left the page, nothing ran, nothing was refused.
    assert b.evaluate("typeof window.__wc_pwned") == "undefined"
    assert _probe(b) == {"flags": [], "parses": [], "created": [], "csp": []}
    assert _requests_since(b, mark, page_url) == []
    assert _csp_since(b, mark) == [] and _errors_since(b, mark) == []
    assert len(s.checked) >= 20 + len(STORED)

    # The CSV: Python's csv module reads it; every value is the stored one;
    # every formula-led value is counted in the note's CAUTION.
    rows = _csv_rows(csv_path)
    assert len(rows) == len(PAYLOADS) + 1
    assert rows[0][0] == PAYLOADS[8]
    assert [r[0] for r in rows[1:]] == STORED
    note = open(note_path, encoding="utf-8-sig").read()
    caution = re.search(r"\n  ([\d,]+) value\(s\) in this file begin with", note)
    assert caution and int(caution.group(1).replace(",", "")) == _csv_formula_cells(rows)
    assert PAYLOADS[20] + " \u2014 " + PAYLOADS[13] in note
    assert "  " + snap.meta["evidence"][0]["path"] + "\n" in note
    filtered = open(filtered_note, encoding="utf-8-sig").read()
    assert "  Search text      : " + PAYLOADS[15].strip() in filtered

    # The workbook: every part parses, openpyxl opens it, no formula
    # anywhere; every value is the stored one; the Provenance sheet holds
    # the hostile labels as text.
    cells, tags, wb = _xlsx_formulas(xlsx_path)
    assert (cells, tags) == (0, 0)
    sheet, prov = wb.worksheets
    got = [sheet.cell(row=r, column=1).value for r in range(2, len(STORED) + 2)]
    assert got == [_xlsx_text(v) for v in STORED]
    assert sheet.cell(row=1, column=1).value == PAYLOADS[8]
    lines = [c.value for c in prov["A"] if c.value is not None]
    assert any(PAYLOADS[13] in v for v in lines)

    # Over the whole walk: no console error at all, and every CSP report
    # came from a positive control (each check asserted none before its
    # plant and at least one after).
    assert b.console_errors() == []


def test_every_instrument_sees_a_planted_injection(probed, tmp_path):
    """Positive controls for the instruments the surface walk cannot plant
    into: the census and CSP log on a test-only page that renders the same
    hostile row through innerHTML; the island/title scan on a page whose
    island was written raw; the canary and the network log on a page with
    no CSP at all; and the two download checkers on planted files."""
    b = probed
    island = island_json(build_island(_hostile_snapshot(), fx.GENERATED))
    # 1. innerHTML of a hostile row, under the page's own CSP.
    script = ("(function () { var d = JSON.parse(document.getElementById("
              "'wc-data').textContent); document.getElementById('app')"
              ".innerHTML = d.views[0].segments[0].rows[%d][0];"
              " window.__wc_ready = true; })();\n" % PAYLOADS.index(PLANT))
    path = tmp_path / "control.html"
    path.write_text(render_page(island, script=script), encoding="utf-8",
                    newline="")
    b.open(str(path))
    _settle(b, 150)
    state = _probe(b)
    assert _census(b)["img"] >= 1
    assert any("data-wcpwn" in f for f in state["flags"])
    assert "innerHTML" in state["parses"]
    assert state["csp"] and b.csp_violations()
    assert b.evaluate("typeof window.__wc_pwned") == "undefined"

    # 2. An island written WITHOUT island_json's escapes breaks out of its
    # element: the scan that found nothing above finds the injection here.
    clean = render_page(island_json(build_island(fx.snapshot(), fx.GENERATED)))
    raw = clean.replace("example.test/page/1\"",
                        "example.test/page/1</script><wcx-pwn data-wcpwn>"
                        "</wcx-pwn><title>x</title><script>\"", 1)
    assert raw != clean
    path = tmp_path / "raw_island.html"
    path.write_text(raw, encoding="utf-8", newline="")
    b.open(str(path), wait_ready=False)
    _settle(b)
    assert b.evaluate("document.querySelectorAll('[data-wcpwn], wcx-pwn').length") == 1
    assert b.evaluate("document.scripts.length") > 2
    assert b.evaluate("document.getElementsByTagName('title').length") == 2

    # 3. No CSP at all: the same payload RUNS (the canary is live) and asks
    # the network for an image (the request log is live).
    page = ("<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>c</title>"
            "</head><body><div id=\"app\"></div><script>"
            "document.getElementById('app').innerHTML = %s;"
            "window.__wc_ready = true;</script></body></html>"
            % json.dumps("<img src=\"http://127.0.0.1:9/wc-probe.png\" "
                         "data-wcpwn onerror=\"%s\">" % CANARY))
    path = tmp_path / "no_csp.html"
    path.write_text(page, encoding="utf-8", newline="")
    b.open(str(path))
    b.evaluate("new Promise(function (ok) { var n = 0; (function wait() {"
               " if (window.__wc_pwned || ++n > 100) ok(); else setTimeout(wait, 50);"
               " })(); })", await_promise=True)
    assert b.evaluate("window.__wc_pwned") == 1
    assert any(u.startswith("http://127.0.0.1:9/") for u in b.requests())
    assert any("data-wcpwn" in f for f in _probe(b)["flags"])
    b.evaluate("void document.createElement('img')")
    assert _probe(b)["created"] == ["img"]

    # 4. The download checkers count a planted formula.
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=1+1"
    planted = str(tmp_path / "planted.xlsx")
    wb.save(planted)
    cells, tags, _wb = _xlsx_formulas(planted)
    assert (cells, tags) == (1, 1)
    planted_csv = tmp_path / "planted.csv"
    planted_csv.write_text("h\r\n=1+1\r\nok\r\n", encoding="utf-8", newline="")
    assert _csv_formula_cells(_csv_rows(str(planted_csv))) == 1
    # And an XML part carrying a character XML 1.0 cannot hold fails the
    # parse the workbook check relies on.
    with pytest.raises(ET.ParseError):
        ET.fromstring("<t>a\x00b</t>".encode("utf-8"))
    os.remove(planted)
