"""Location shows the raw records behind its fixes, and every position
export carries each fix's source record."""
import csv
import importlib
import json

import openpyxl
import pytest

from webcache.engine.forensic_engine import ENTRY_FIELDS, ForensicEngine
from webcache.provenance import CARVED_ORIGIN
from webcache.shell import location_loader
from webcache.views.location import LocationView

RAW = 133544180420000000


def _geo_entry(i, query="lat=51.5&lon=-0.12"):
    url = "https://example.test/geo?%s&n=%d" % (query, i)
    return {"EntryId": i, "ContainerId": 4, "ContainerType": "History",
            "Url": url, "CleanUrl": url, "AccessedTime_raw": RAW + i,
            "ModifiedTime": "2024-03-11 09:14:02 UTC"}


def _carved():
    url = "https://example.test/geo?lat=48.8566&lon=2.3522"
    return {"Url": url, "CleanUrl": url, "Source": "V01.log", "Offset": 4096,
            "RecoverySource": "Structured record reconstruction",
            "Confidence": "high", "ContainerType": "History"}


def test_the_map_is_gone(qtbot):
    with pytest.raises(ModuleNotFoundError):
        importlib.import_module("webcache.widgets.fix_plot")
    view = LocationView()
    qtbot.addWidget(view)
    assert not hasattr(view, "plot")


def test_location_has_three_table_modes(qtbot):
    view = LocationView()
    qtbot.addWidget(view)
    assert [(rs.key, rs.label) for rs in view._record_sets] == [
        ("fixes", "Position fixes"), ("places", "Places"),
        ("records", "Raw records")]


def test_one_row_per_source_record_with_its_canonical_fields():
    fixes = location_loader.position_fixes([_geo_entry(1), _geo_entry(2)])
    rows = location_loader.source_records(fixes)
    assert sorted(r["EntryId"] for r in rows) == [1, 2]
    for row in rows:
        assert row["Fixes"] == 1
        assert row["AccessedTime_raw"] == RAW + row["EntryId"]
        assert any(row["_record"] is f["_record"] for f in fixes)
        # Derived by the engine's own row builder, not copied by hand.
        for key in ("TableName", "Domain", "KeyTimestampUTC",
                    "AccessedTime_rfc"):
            assert key in row, key


def test_the_raw_records_grid_has_a_column_for_every_canonical_field(qtbot):
    view = LocationView()
    qtbot.addWidget(view)
    sets = {rs.key: rs for rs in view._record_sets}
    assert "records" in sets, "no Raw records mode"
    records = sets["records"]
    keys = [c.key for c in records.columns]
    assert keys[0] == "Fixes" and keys[1] == "Origin"
    assert sorted(keys[1:]) == sorted(k for k, _h, _g in ENTRY_FIELDS)


def test_a_record_with_two_fixes_is_one_row_counting_both():
    fixes = location_loader.position_fixes(
        [_geo_entry(1, "ll=48.8566,2.3522&lat=51.5&lon=-0.12")])
    assert len(fixes) == 2
    rows = location_loader.source_records(fixes)
    assert len(rows) == 1 and rows[0]["Fixes"] == 2


def test_carved_records_are_labelled():
    fixes = location_loader.position_fixes([], [_carved()])
    rows = location_loader.source_records(fixes)
    assert rows[0]["Origin"] == CARVED_ORIGIN
    assert rows[0]["Offset"] == 4096


def _fixes():
    return location_loader.position_fixes([_geo_entry(1)], [_carved()])


def _expected_header():
    return ([h for _k, h in ForensicEngine.LOCATION_FIELDS]
            + [ForensicEngine.LOCATION_RECORD_PREFIX + h
               for _k, h, _g in ENTRY_FIELDS])


def test_the_location_csv_writes_the_fix_then_its_source_record(tmp_path):
    path = str(tmp_path / "fixes.csv")
    ForensicEngine().export_location_csv(_fixes(), path)
    with open(path, encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    assert rows[0] == _expected_header()
    col = rows[0].index("Source record: AccessedTime (FILETIME)")
    assert str(RAW + 1) in [r[col] for r in rows[1:]]
    offset = rows[0].index("Source record: Carved at offset")
    assert "4096" in [r[offset] for r in rows[1:]]


def test_the_location_json_writes_the_fix_and_its_source_record(tmp_path):
    path = str(tmp_path / "fixes.json")
    assert ForensicEngine().export_location_json(_fixes(), path) == 2
    data = json.load(open(path, encoding="utf-8"))
    first = data["fixes"][0]
    assert list(first["fix"]) == [k for k, _h in ForensicEngine.LOCATION_FIELDS]
    assert list(first["source_record"]) == [k for k, _h, _g in ENTRY_FIELDS]


def test_the_location_xlsx_matches_the_csv_layout(tmp_path):
    path = str(tmp_path / "fixes.xlsx")
    assert ForensicEngine().export_location_xlsx(_fixes(), path) == 2
    ws = openpyxl.load_workbook(path).active
    assert [c.value for c in ws[1]] == _expected_header()


def test_the_kml_stays_points(tmp_path):
    path = str(tmp_path / "fixes.kml")
    ForensicEngine().export_location_kml(_fixes(), path)
    text = open(path, encoding="utf-8").read()
    assert text.count("<Point>") == 2
    # The fix's own "Source record: Live record" line is one of
    # LOCATION_FIELDS and belongs there; what must NOT appear is the
    # appended canonical block the tabular formats carry.
    for _k, header, _g in ENTRY_FIELDS:
        assert "Source record: %s:" % header not in text, header


# --- --------------------------------------------------------

def test_the_raw_records_origin_column_is_headed_origin_detail(qtbot):
    """The vertical header already reads "Origin"."""
    view = LocationView()
    qtbot.addWidget(view)
    records = {rs.key: rs for rs in view._record_sets}["records"]
    labels = {c.key: c.label for c in records.columns}
    assert labels["Origin"] == "Origin detail"


def _surrogate_carved():
    url = "https://example.test/geo?lat=48.8566&lon=2.3522&q=\udcff\x01"
    return {"Url": url, "CleanUrl": url, "Source": "V01.log", "Offset": 8192,
            "RecoverySource": "Structured record reconstruction",
            "Confidence": "high", "ContainerType": "History"}


def test_the_location_json_survives_a_lone_surrogate(tmp_path):
    """Carved bytes regularly decode to lone surrogates (_safe_text's own
    docstring); the JSON must still be written, as the CSV and xlsx are."""
    fixes = location_loader.position_fixes([], [_surrogate_carved()])
    assert fixes
    path = str(tmp_path / "fixes.json")
    assert ForensicEngine().export_location_json(fixes, path) == len(fixes)
    data = json.load(open(path, encoding="utf-8"))
    assert len(data["fixes"]) == len(fixes)




def test_every_location_time_column_holds_dates(tmp_path):
    """Exported from a real exhibit, the position workbook's "Fix time --
    device said (UTC)" and "Cache wrote (UTC)" were TEXT: only the source
    record's columns were flagged as timestamps, so the two times the
    workbook exists for sorted alphabetically. Same rule as the forensic
    workbook: a heading that names a clock holds dates. A fix time the
    device sent in a shape that is not a date is left as it was sent."""
    fixes = _fixes()
    fixes[0]["FixTimestamp"] = "2026-02-17T16:50:09Z"
    fixes[1]["FixTimestamp"] = "not a time"
    path = str(tmp_path / "fixes.xlsx")
    ForensicEngine().export_location_xlsx(fixes, path)
    ws = openpyxl.load_workbook(path).active
    header = [c.value for c in ws[1]]
    fix = [ws.cell(row=r, column=header.index(
        "Fix time — device said (UTC)") + 1).value for r in (2, 3)]
    import datetime
    assert isinstance(fix[0], datetime.datetime)
    assert fix[1] == "not a time"
    wrote = header.index("Cache wrote (UTC)") + 1
    for r in (2, 3):
        value = ws.cell(row=r, column=wrote).value
        assert value is None or isinstance(value, datetime.datetime), value
