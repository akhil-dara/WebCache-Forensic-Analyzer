"""The page's workbook is the app's view workbook, cell for cell (spec:
"Export from the page"). Read back with openpyxl, both files."""
import datetime
import zipfile

import openpyxl
import pytest

from webcache.shell import export_formats

from tests import _package_fixtures as fx

STATEMENT = ("[TRUNCATED BY EXCEL: 40,000 characters total, 8,000 not shown. "
             "The complete value is in the CSV and JSON exports and in the "
             "app's blob viewer.]")

#: Values that would break a workbook written by string concatenation, or
#: run as a formula when it is opened. Synthetic, every one.
HOSTILE = [
    '</t></is></c><c r="Z9"><f>1+1</f></c><c><is><t>',
    "a & b < c > d \"e\" 'f'",
    "]]><![CDATA[x]]>",
    "line one\r\nline two\rline three\n",
    "\ttabbed\t",
    "  padded  ",
    "+1+2",
    "-3",
    "@SUM(A1:A2)",
    "=1+1",
    "&#0; &amp; &lt;",
    "_x0041_ stays as typed",
    "smile \U0001F600 and é",
]


def _rows():
    rows = ([fx.history_row(i) for i in range(1, 4)]
            + [fx.history_row(i, carved=True) for i in (7, 8)])
    rows[0]["CleanUrl"] = '=HYPERLINK("https://example.test")'
    rows[1]["CleanUrl"] = "2026-02-17 16:50:09"
    rows[2]["_record"]["UrlHash"] = 2 ** 60
    rows[3]["CleanUrl"] = "x" * 40000
    return rows


def _cells(ws):
    return [[(c.value, c.data_type, c.number_format, bool(c.font and c.font.b))
             for c in row] for row in ws.iter_rows()]


def _widths(ws):
    """Each column's width, however the writer grouped its <col> ranges
    (openpyxl merges equal neighbours into one min..max range)."""
    out = {}
    for dim in ws.column_dimensions.values():
        if dim.min and dim.max:
            for i in range(dim.min, dim.max + 1):
                out[i] = dim.width
    return [out.get(i) for i in range(1, ws.max_column + 1)]


def _formulas(wb):
    return sum(1 for ws in wb.worksheets for row in ws.iter_rows()
               for c in row if c.data_type == "f")


def _export(browser, tmp_path, rows, store, before=None):
    browser.clear_downloads()
    browser.open(fx.write(tmp_path, fx.snapshot(views=[fx.history_view(
        [fx.segment(rows=rows)])])))
    if before:
        browser.evaluate(before)
    browser.evaluate("WC.exportXlsx('history', 'history', %s)"
                     % ("true" if store else "false"), await_promise=True)
    return browser.wait_download("webcache_history.xlsx")


def _app(tmp_path, rows):
    app_path = tmp_path / "app.xlsx"
    export_formats.write_view_xlsx(rows, fx.history_columns(), str(app_path),
                                   {"view_label": "History"})
    return openpyxl.load_workbook(app_path)


def _methods(path):
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        return {info.compress_type for info in z.infolist()}


@pytest.mark.parametrize("store", [False, True], ids=["deflated", "stored"])
def test_the_workbook_is_the_apps_view_workbook_cell_for_cell(browser,
                                                              tmp_path, store):
    rows = _rows()
    path = _export(browser, tmp_path, rows, store)
    assert _methods(path) == ({zipfile.ZIP_STORED} if store
                              else {zipfile.ZIP_DEFLATED})
    page = openpyxl.load_workbook(path)
    app = _app(tmp_path, rows)
    assert page.sheetnames == ["History", "Provenance"]
    ws, ref = page["History"], app["History"]
    assert _cells(ws) == _cells(ref)
    assert ws.freeze_panes == ref.freeze_panes == "A2"
    assert ws.auto_filter.ref == ref.auto_filter.ref
    assert _widths(ws) == _widths(ref)
    assert page["Provenance"]["A1"].value == "Provenance of sheet 1"
    assert page["Provenance"]["A2"].value.startswith(
        "WebCache Forensic Analyzer")
    assert _formulas(page) == 0


def test_a_browser_without_compression_streams_writes_a_stored_workbook(
        browser, tmp_path):
    rows = _rows()
    path = _export(browser, tmp_path, rows, False,
                   before="window.CompressionStream = undefined; true")
    assert _methods(path) == {zipfile.ZIP_STORED}
    page = openpyxl.load_workbook(path)
    assert _cells(page["History"]) == _cells(_app(tmp_path, rows)["History"])


def test_only_timestamp_columns_hold_dates_and_text_stays_text(browser,
                                                               tmp_path):
    rows = _rows()
    ws = openpyxl.load_workbook(_export(browser, tmp_path, rows, False))["History"]
    headers = [c.value for c in ws[1]]
    address = headers.index("Address") + 1
    visit = headers.index("Last visit (UTC)") + 1
    url_hash = headers.index("UrlHash") + 1
    assert ws.cell(row=2, column=address).value == '=HYPERLINK("https://example.test")'
    assert ws.cell(row=2, column=address).data_type == "s"
    assert ws.cell(row=3, column=address).value == "2026-02-17 16:50:09"
    assert isinstance(ws.cell(row=2, column=visit).value, datetime.datetime)
    assert ws.cell(row=4, column=url_hash).value == str(2 ** 60)
    long_cell = ws.cell(row=5, column=address).value
    assert long_cell == "x" * 32000 + "…" + STATEMENT


@pytest.mark.parametrize("store", [False, True], ids=["deflated", "stored"])
def test_hostile_values_make_a_valid_workbook_with_no_formulas(browser,
                                                               tmp_path, store):
    rows = [fx.history_row(i) for i in range(1, len(HOSTILE) + 1)]
    for row, value in zip(rows, HOSTILE):
        row["CleanUrl"] = value
        row["_record"]["Url"] = value
    page = openpyxl.load_workbook(_export(browser, tmp_path, rows, store))
    assert _formulas(page) == 0
    assert _cells(page["History"]) == _cells(_app(tmp_path, rows)["History"])
    address = [c.value for c in page["History"][1]].index("Address")
    assert [r[address].value for r in
            page["History"].iter_rows(min_row=2)] == HOSTILE


def test_a_character_xml_cannot_hold_is_written_as_a_visible_escape(
        browser, tmp_path):
    rows = [fx.history_row(1)]
    rows[0]["CleanUrl"] = "a￾b￿c"
    ws = openpyxl.load_workbook(_export(browser, tmp_path, rows, False))["History"]
    address = [c.value for c in ws[1]].index("Address") + 1
    assert ws.cell(row=2, column=address).value == "a\\uFFFEb\\uFFFFc"


def test_a_datetime_outside_a_timestamp_column_states_its_true_instant(
        browser, tmp_path):
    """Outside a timestamp column a datetime is text. An
    aware one is converted to UTC before it gets its "Z"; a naive one gets
    no "Z" (nothing says it is UTC); microseconds are kept. The page writes
    the app's text cell for cell."""
    from webcache.engine import forensic_engine
    ist = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
    aware = datetime.datetime(2026, 2, 17, 21, 20, 9, 123456, tzinfo=ist)
    naive = datetime.datetime(2026, 2, 17, 21, 20, 9)
    assert forensic_engine._cell(aware) == "2026-02-17T15:50:09.123456Z"
    assert forensic_engine._cell(naive) == "2026-02-17T21:20:09"
    from webcache.shell.html_package import PackageColumn
    from webcache.widgets.column_spec import Column
    note = Column("Note", "Note", 160, kind="text", field="Note")
    cols = fx.history_columns() + [PackageColumn.from_column(note)]
    rows = [fx.history_row(1, Note=aware), fx.history_row(2, Note=naive),
            fx.history_row(3)]
    browser.clear_downloads()
    browser.open(fx.write(tmp_path, fx.snapshot(views=[fx.history_view(
        [fx.segment(rows=rows, columns=cols)])])))
    browser.evaluate("WC.exportXlsx('history', 'history', false)",
                     await_promise=True)
    page = openpyxl.load_workbook(browser.wait_download("webcache_history.xlsx"))
    app_path = tmp_path / "app.xlsx"
    export_formats.write_view_xlsx(rows, cols, str(app_path),
                                   {"view_label": "History"})
    app = openpyxl.load_workbook(app_path)
    assert _cells(page["History"]) == _cells(app["History"])
    texts = [c.value for row in app["History"].iter_rows() for c in row]
    assert "2026-02-17T15:50:09.123456Z" in str(texts)
    assert "2026-02-17T21:20:09" in str(texts)
