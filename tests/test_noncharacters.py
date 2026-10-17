"""U+FFFE and U+FFFF in evidence: every writer must write them as the same
visible escape the offline page uses, never refuse the file.
openpyxl rejects both characters, so before this fix the app's XLSX export
wrote nothing at all for a record that carried one."""
import json

from openpyxl import load_workbook

from webcache.engine import forensic_engine
from webcache.shell import export_formats, html_package

VALUE = "a￾b￿c"
ESCAPED = "a\\uFFFEb\\uFFFFc"


def test_safe_text_escapes_both_noncharacters():
    assert forensic_engine._safe_text(VALUE) == ESCAPED


def test_the_app_xlsx_is_written_and_reads_back_the_escape(tmp_path):
    path = tmp_path / "v.xlsx"
    forensic_engine.write_simple_xlsx(str(path), "View", ["Address"],
                                      [[VALUE]])
    ws = load_workbook(path).active
    assert ws.cell(row=2, column=1).value == ESCAPED


def test_the_island_reads_back_the_same_escape():
    text = html_package.island_json({"label": VALUE})
    assert json.loads(text)["label"] == ESCAPED


def test_the_encoding_paragraph_names_both_noncharacters():
    import inspect
    source = inspect.getsource(export_formats)
    assert "U+FFFE" in source and "U+FFFF" in source
    page = (html_package._ASSET_DIR / "export_csv.js").read_text("utf-8")
    assert "U+FFFE" in page and "U+FFFF" in page
