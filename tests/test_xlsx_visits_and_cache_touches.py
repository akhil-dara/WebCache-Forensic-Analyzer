"""The workbook's focused sheets keep visits and cache touches apart.

Its History "Visits" and Downloads "Opens" columns held VisitCountBest:
the property-store count where there was one, otherwise AccessCount
(cache touches), with nothing in the cell to say which. The grids and
the PDF already give each counter its own column."""
import openpyxl

from webcache.engine import forensic_engine

BASE = {"ContainerType": "History", "ContainerName": "History",
        "ContainerId": 1, "AccessedTime_raw": 133843382090000000}


def _sheet(tmp_path, name, entries):
    engine = forensic_engine.ForensicEngine()
    engine.load_entries(entries)
    path = str(tmp_path / "b.xlsx")
    engine.export_xlsx(entries, path, True, {"examiner": "E"}, {}, {},
                       None, None, None, engine.get_statistics())
    ws = openpyxl.load_workbook(path)[name]
    heads = [c.value for c in ws[1]]
    rows = [{h: c.value for h, c in zip(heads, r)}
            for r in ws.iter_rows(min_row=2)]
    return heads, rows


def test_history_has_a_column_for_each_counter(tmp_path):
    heads, rows = _sheet(tmp_path, "History", [
        dict(BASE, EntryId=1, Url="https://site.test/a", CleanUrl="https://site.test/a", SPS_VisitCount=3,
             AccessCount=9),
        dict(BASE, EntryId=2, Url="https://site.test/b", CleanUrl="https://site.test/b", AccessCount=7)])
    assert "Visits" not in heads
    assert {"Visits (property store)", "Cache touches"} <= set(heads)
    by_url = {r["URL"]: r for r in rows}
    a, b = by_url["https://site.test/a"], by_url["https://site.test/b"]
    assert (a["Visits (property store)"], a["Cache touches"]) == (3, 9)
    assert b["Visits (property store)"] in (None, "")
    assert b["Cache touches"] == 7


def test_downloads_has_a_column_for_each_counter(tmp_path):
    heads, rows = _sheet(tmp_path, "Downloads & Files", [
        dict(BASE, EntryId=1, Url="file:///C:/x/report.pdf", IsLocalFile=True,
             AccessCount=4)])
    assert "Opens" not in heads
    assert {"Opens (property store)", "Cache touches"} <= set(heads)
    assert rows[0]["Cache touches"] == 4
    assert rows[0]["Opens (property store)"] in (None, "")
