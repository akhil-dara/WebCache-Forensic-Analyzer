"""A coordinate is shown and exported exactly as the URL carried it.

Synthetic: trailing zeros, and 18 significant digits -- more than a double
holds, so any float round trip changes the string."""
import csv
import json

import openpyxl
from PySide6.QtCore import Qt

from webcache.engine import geolocation
from webcache.engine.forensic_engine import ForensicEngine
from webcache.shell import location_loader
from webcache.views.location import LocationView

LAT = "12.345678901234567800"
LON = "-98.765432109876543200"
URL = "https://example.test/weather?lat=%s&lon=%s" % (LAT, LON)
PAIR_URL = "https://example.test/feed?ll=%s%%2C%s" % (LAT, LON)
ENTRY = {"EntryId": 1, "ContainerId": 4, "ContainerType": "History",
         "Url": URL, "CleanUrl": URL, "ModifiedTime": "2024-03-11 09:14:02 UTC"}


def _fixes():
    return location_loader.position_fixes([ENTRY])


def test_the_strings_survive_extraction():
    fix = geolocation.extract_from_url(URL)[0]
    assert (fix["LatitudeRaw"], fix["LongitudeRaw"]) == (LAT, LON)
    assert fix["Latitude"] == float(LAT)          # still a float, for sorting


def test_a_pair_parameter_keeps_its_strings_after_percent_decoding():
    fix = geolocation.extract_from_url(PAIR_URL)[0]
    assert (fix["LatitudeRaw"], fix["LongitudeRaw"]) == (LAT, LON)


def test_the_map_links_embed_the_strings():
    fix = geolocation.extract_from_url(URL)[0]
    assert fix["MapsUrl"] == "https://www.google.com/maps?q=%s,%s" % (LAT, LON)
    assert "mlat=%s&mlon=%s" % (LAT, LON) in fix["OsmUrl"]


def test_the_view_model_shows_the_strings(qtbot):
    view = LocationView()
    qtbot.addWidget(view)
    view.set_records_for("fixes", _fixes())
    model = view.table.model()
    col = [c.key for c in model.columns].index("Coords")
    assert model.data(model.index(0, col), Qt.DisplayRole) == "%s, %s" % (LAT, LON)


def test_the_location_csv_holds_the_strings(tmp_path):
    path = str(tmp_path / "fixes.csv")
    ForensicEngine().export_location_csv(_fixes(), path)
    with open(path, encoding="utf-8-sig", newline="") as fh:
        header, row = list(csv.reader(fh))[:2]
    assert row[header.index("Latitude (as stored)")] == LAT
    assert row[header.index("Longitude (as stored)")] == LON


def test_the_kml_coordinates_are_the_strings(tmp_path):
    path = str(tmp_path / "fixes.kml")
    ForensicEngine().export_location_kml(_fixes(), path)
    text = open(path, encoding="utf-8").read()
    assert "<coordinates>%s,%s,0</coordinates>" % (LON, LAT) in text


def test_the_json_holds_the_strings(tmp_path):
    path = str(tmp_path / "fixes.json")
    ForensicEngine().export_location_json(_fixes(), path)
    fix = json.load(open(path, encoding="utf-8"))["fixes"][0]["fix"]
    assert (fix["LatitudeRaw"], fix["LongitudeRaw"]) == (LAT, LON)


def test_the_xlsx_holds_the_strings_as_text(tmp_path):
    path = str(tmp_path / "fixes.xlsx")
    ForensicEngine().export_location_xlsx(_fixes(), path)
    ws = openpyxl.load_workbook(path).active
    header = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=header.index("Latitude (as stored)") + 1).value == LAT
    assert ws.cell(row=2, column=header.index("Longitude (as stored)") + 1).value == LON


def test_a_place_keeps_its_own_labelled_rounding():
    places = location_loader.places(_fixes(), grouping="11m")
    assert places[0]["MapsUrl"].endswith("q=12.3457,-98.7654")
