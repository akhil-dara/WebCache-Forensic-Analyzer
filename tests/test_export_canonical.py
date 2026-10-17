"""A view export writes the view's columns, then every canonical field
its records carry -- raw FILETIMEs, ids and carve provenance included."""
import csv

import openpyxl

from webcache.engine import forensic_engine
from webcache.provenance import CARVED_ORIGIN
from webcache.shell import export_formats as F
from webcache.shell.record_kinds import K_CARVED_STRUCT, K_HISTORY
from webcache.widgets.column_spec import Column
from tests.test_export_runner import _job, _run

RAW = 133544180420000001

LIVE = {"EntryId": 7, "ContainerId": 4, "ContainerType": "History",
        "Url": "https://example.test/a", "CleanUrl": "https://example.test/a",
        "UrlHash": 99, "AccessedTime_raw": RAW,
        "ModifiedTime": "2024-03-11 09:14:02 UTC", "ModifiedTime_raw": RAW - 1}
CARVED = {"Url": "https://example.test/b", "CleanUrl": "https://example.test/b",
          "RecoverySource": "Structured record reconstruction",
          "Source": "V01.log", "Offset": 4096, "Page": 3,
          "Confidence": "high", "ContainerId": None}

COLS = [Column("CleanUrl", "Address", 400, kind="url")]


def _rows():
    return [
        {"CleanUrl": LIVE["CleanUrl"], "Origin": "Live record",
         "_record": LIVE, "_kind": K_HISTORY},
        {"CleanUrl": CARVED["CleanUrl"], "Origin": CARVED_ORIGIN,
         "Source": "V01.log", "_record": CARVED, "_kind": K_CARVED_STRUCT},
    ]


def _read(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.reader(fh))


def test_canonical_record_row_is_the_engines_own_row():
    assert forensic_engine.canonical_record_row(LIVE) == \
        forensic_engine.build_entry_row(LIVE)
    assert forensic_engine.canonical_record_row(CARVED) == \
        forensic_engine.build_recovered_row(CARVED)


def test_the_views_columns_come_first_then_every_canonical_field(tmp_path):
    path = str(tmp_path / "v.csv")
    F.write_view_csv(_rows(), COLS, path, {})
    header, live, carved = _read(path)
    assert header[:3] == ["Address", "Origin", "RecordSource"]
    for key, head, _g in forensic_engine.ENTRY_FIELDS:
        if key in ("CleanUrl", "Origin", "Source"):
            continue
        assert head in header, head
    assert live[header.index("AccessedTime (FILETIME)")] == str(RAW)
    assert live[header.index("EntryId")] == "7"
    assert live[header.index("UrlHash")] == "99"
    assert carved[header.index("Carved at offset")] == "4096"
    assert carved[header.index("Carve confidence")] == "high"


def test_derived_columns_are_never_dropped(tmp_path):
    path = str(tmp_path / "v.csv")
    F.write_view_csv(_rows(), COLS, path, {})
    header = _read(path)[0]
    for derived in ("ModifiedTime (UTC)", "ModifiedTime (RFC3339)",
                    "KeyTimestamp (RFC3339 UTC)", "Domain"):
        assert derived in header, derived


def test_rows_without_a_record_add_no_canonical_columns(tmp_path):
    path = str(tmp_path / "v.csv")
    F.write_view_csv([{"CleanUrl": "x", "_kind": "domain"}], COLS, path, {})
    assert _read(path)[0] == ["Address", "Origin", "RecordSource"]


def test_the_xlsx_carries_the_same_columns_and_raw_values(tmp_path):
    path = str(tmp_path / "v.xlsx")
    n = F.write_view_xlsx(_rows(), COLS, path, {"view_label": "History"})
    assert n == 2
    ws = openpyxl.load_workbook(path).active
    header = [c.value for c in ws[1]]
    csv_path = str(tmp_path / "v.csv")
    F.write_view_csv(_rows(), COLS, csv_path, {})
    assert header == _read(csv_path)[0]
    cell = ws.cell(row=2, column=header.index("AccessedTime (FILETIME)") + 1)
    assert str(cell.value) == str(RAW), "a FILETIME above 2**53 must not be rounded"


def test_the_sidecar_states_both_column_counts(tmp_path):
    from webcache.widgets import record_filter

    class _Model:
        columns = COLS
        filter = record_filter.RecordFilter(_rows(), COLS)

        def rowCount(self):
            return self.filter.visible_count

        def chip_text(self):
            return ""

    scope = F.scope_from_view(object(), "history", "History", _Model())
    canon = F.canonical_columns_for(_rows(), COLS)
    assert scope["view_column_count"] == 3
    assert scope["canonical_column_count"] == len(canon) > 40
    path = str(tmp_path / "v.csv")
    F.write_view_csv(_rows(), COLS, path, scope)
    text = open(F.write_view_provenance(path, {}, scope, 0),
                encoding="utf-8-sig").read()
    assert "View columns     : 3" in text
    assert "Canonical fields : %d" % len(canon) in text


def test_the_runner_writes_a_view_xlsx(tmp_path):
    path = str(tmp_path / "history.xlsx")
    out = _run(_job("view_xlsx", path, rows=_rows(), columns=COLS,
                    scope={"view": "history", "view_label": "History",
                           "columns": [], "hidden_columns": [],
                           "rows_in_file": 2, "rows_in_view": 2,
                           "rows_in_record_set": 2, "is_subset": False,
                           "text_filter": "", "conditions": {}, "values": {},
                           "chip_text": "", "live_rows": 1, "carved_rows": 1,
                           "recovery_scan_run": True},
                    stamp={"exported_utc": "2026-09-26T00:00:00Z",
                           "evidence": {}}))
    assert out["failed"] == []
    assert path in out["written"]


def test_export_view_offers_xlsx(qtbot, tmp_path, monkeypatch):
    from webcache.shell.main_window import MainWindow
    w = MainWindow()
    qtbot.addWidget(w)
    view = w._ensure_view_built("history")
    view.set_records(_rows())
    jobs = []
    monkeypatch.setattr(w, "_choose_export_path",
                        lambda *a: str(tmp_path / "history.xlsx"))
    monkeypatch.setattr(w, "_start_export",
                        lambda job, card_key=None: jobs.append(job))
    w._export_view(view, "history")
    assert [j.key for j in jobs] == ["view_xlsx"]


def _file_access_rows():
    """A File Access row set built by the app's own loader, so the export
    sees exactly the row shape the screen holds."""
    from webcache.shell.evidence_loader import file_access_rows_from_parser

    class _Parser:
        all_entries = [{
            "EntryId": 11, "ContainerId": 5, "ContainerType": "History",
            "Url": "file:///C:/synthetic/report.txt",
            "CleanUrl": "file:///C:/synthetic/report.txt",
            "IsLocalFile": True, "UrlHash": RAW + 7,
            "CreationTime_raw": RAW + 2, "ModifiedTime_raw": RAW + 3,
            "AccessedTime_raw": RAW + 4, "ExpiryTime_raw": RAW + 5,
            "SyncTime_raw": RAW + 6, "PostCheckTime_raw": RAW + 8,
            "SPS_LastVisitLocal_raw": RAW + 9,
        }]
    rows = file_access_rows_from_parser(_Parser())
    assert len(rows) == 1
    return rows, _Parser.all_entries[0]


_RAW_HEADS = {k: h for k, h, _g in forensic_engine.ENTRY_FIELDS
              if k.endswith("_raw")}


def test_file_access_export_carries_every_raw_filetime_exactly(tmp_path):
    """The examiner's words: 'raw filetime etc properly and not derived
    only'. Every *_raw value is above 2**53, where a float would round it;
    both formats must hold the parser's integer to the last digit."""
    from webcache.views.file_access import FILE_ACCESS_COLUMNS
    rows, entry = _file_access_rows()
    csv_path = str(tmp_path / "fa.csv")
    F.write_view_csv(rows, FILE_ACCESS_COLUMNS, csv_path, {})
    header, line = _read(csv_path)
    assert header[:len(FILE_ACCESS_COLUMNS)] == \
        [c.label for c in FILE_ACCESS_COLUMNS]
    for key, head in _RAW_HEADS.items():
        assert entry[key] > 2 ** 53
        assert head in header, head
        assert line[header.index(head)] == str(entry[key]), head
    xlsx_path = str(tmp_path / "fa.xlsx")
    F.write_view_xlsx(rows, FILE_ACCESS_COLUMNS, xlsx_path,
                      {"view_label": "File Access"})
    ws = openpyxl.load_workbook(xlsx_path).active
    xheader = [c.value for c in ws[1]]
    assert xheader == header
    for key, head in _RAW_HEADS.items():
        cell = ws.cell(row=2, column=xheader.index(head) + 1)
        assert str(cell.value) == str(entry[key]), head


# --- --------------------------------------------------------

_URL_ONLY_HEADS = ("UrlCategory", "AccessedTime (FILETIME)", "Scheme",
                   "KeyTimestamp (RFC3339 UTC)")


def _cookie_rows():
    from webcache.shell.evidence_loader import cookie_rows_from_parser

    class _Parser:
        cookie_entries = [{
            "EntryId": 21, "Table": "CookieEntryEx_3", "PartitionId": "p0",
            "Domain": "example.test", "RDomain": "test.example",
            "Name": "sid", "Value": "abc", "Path": "/", "Flags": 1,
            "IsSecure": True, "Expires": "2025-01-01 00:00:00",
            "Expires_raw": RAW + 11, "LastModified": "2024-01-01 00:00:00",
            "LastModified_raw": RAW + 12,
        }]
    return cookie_rows_from_parser(_Parser()), _Parser.cookie_entries[0]


def _container_rows():
    from webcache.shell.evidence_loader import container_rows_from_parser

    class _Parser:
        containers_meta = [{
            "ContainerId": 4, "Name": "History", "HasTable": True,
            "RecordCount": 3, "Size": 10, "PartitionId": "p0",
            "Directory": "C:/synthetic/History/", "Flags": 9,
            "LastAccessTime_raw": RAW + 13, "LastScavengeTime_raw": RAW + 14,
        }]

        def get_container_type(self, cid):
            return "History"
    return container_rows_from_parser(_Parser()), _Parser.containers_meta[0]


def _cookie_columns():
    from webcache.views.cookies import COOKIE_COLUMNS
    return COOKIE_COLUMNS


def test_a_cookie_export_carries_its_own_raw_fields_not_the_url_schema(tmp_path):
    """I1: a cookie is not a URL-cache entry. It gets every key its own
    record carries (its raw FILETIMEs included), never ENTRY_FIELDS."""
    rows, rec = _cookie_rows()
    path = str(tmp_path / "c.csv")
    F.write_view_csv(rows, _cookie_columns(), path, {})
    header, line = _read(path)
    for head in _URL_ONLY_HEADS:
        assert head not in header, head
    assert line[header.index("Expires_raw")] == str(RAW + 11)
    assert line[header.index("LastModified_raw")] == str(RAW + 12)
    for key in ("EntryId", "Table", "PartitionId", "RDomain", "Flags"):
        assert line[header.index(key)] == str(rec[key]), key


def test_a_container_export_carries_its_own_raw_fields(tmp_path):
    from webcache.widgets.column_spec import Column as C
    rows, rec = _container_rows()
    cols = [C("Id", "ID", 60, kind="num"), C("Name", "Name", 200),
            C("LastAccess", "Last access (UTC)", 165, kind="time")]
    path = str(tmp_path / "k.xlsx")
    F.write_view_xlsx(rows, cols, path, {"view_label": "Containers"})
    ws = openpyxl.load_workbook(path).active
    header = [c.value for c in ws[1]]
    for head in _URL_ONLY_HEADS + ("Flags (int)", "SourceTable"):
        assert head not in header, head
    for key in ("LastAccessTime_raw", "LastScavengeTime_raw", "Flags",
                "Directory", "HasTable"):
        assert key in header, key
    assert str(ws.cell(row=2, column=header.index(
        "LastScavengeTime_raw") + 1).value) == str(RAW + 14)


IOC_COLS = [Column("Reasons", "Reasons", 220), Column("Source", "Source", 70,
                                                       kind="flag"),
            Column("CleanUrl", "URL", 380, kind="url")]


def _carved_ioc_hit():
    rec = dict(CARVED, Offset=8192, Source="V01.log",
               RecoverySource="Transaction log")
    return {"Reasons": "watchlist", "Source": "carved",
            "CleanUrl": rec["CleanUrl"], "Origin": CARVED_ORIGIN,
            "_record": rec, "_kind": K_CARVED_STRUCT}


def test_a_view_key_that_means_something_else_keeps_the_canonical_field(tmp_path):
    """I2: the IOC view's own Source ("carved"/"live") must not shadow the
    canonical Source -- the file the record was carved from."""
    path = str(tmp_path / "ioc.csv")
    F.write_view_csv([_carved_ioc_hit()], IOC_COLS, path, {})
    header, line = _read(path)
    assert line[header.index("Source")] == "carved"
    assert "Carved from file" in header
    assert line[header.index("Carved from file")] == "V01.log"


def test_a_carved_ioc_hit_states_it_is_carved(tmp_path):
    """I3: ioc_hit_rows set no Origin, so a carved hit's Origin cell said
    "Live record" beside its own carve method. The loader states it."""
    from webcache.shell.iocs_loader import ioc_hit_rows
    live = [dict(LIVE, Url="https://watch.example.test/live",
                 CleanUrl="https://watch.example.test/live")]
    carved = [dict(CARVED, Url="https://watch.example.test/carved",
                   CleanUrl="https://watch.example.test/carved",
                   RecoverySource="Transaction log")]
    rows = ioc_hit_rows(live, carved, terms=["watch.example.test"])
    by_url = {r["CleanUrl"]: r for r in rows}
    assert by_url["https://watch.example.test/carved"].get("Origin") == CARVED_ORIGIN
    assert by_url["https://watch.example.test/live"].get("Origin") == "Live record"
    path = str(tmp_path / "ioc.csv")
    F.write_view_csv(rows, IOC_COLS, path, {})
    header, *lines = _read(path)
    origin = {line[header.index("URL")]: line[header.index("Origin")]
              for line in lines}
    assert origin["https://watch.example.test/carved"] == CARVED_ORIGIN


def _stored_serial(path, ref):
    """The number the sheet XML actually holds for one cell. openpyxl's
    reader rounds a date to the millisecond; the file keeps microseconds."""
    import re
    import zipfile
    xml = zipfile.ZipFile(path).read("xl/worksheets/sheet1.xml").decode()
    m = re.search(r'<c r="%s"[^>]*>(?:<f>.*?</f>)?<v>([^<]*)</v>' % ref, xml)
    return m.group(1) if m else None


def test_view_time_columns_are_real_dates_at_full_precision(tmp_path):
    """I4: a tz-aware view timestamp came out of the xlsx as TEXT with the
    sub-seconds cut off, while the CSV kept them. Both formats must hold
    the same instant to the microsecond, the xlsx as a real date cell."""
    import datetime as dt
    from openpyxl.utils import get_column_letter
    utc = dt.datetime(2026, 8, 3, 15, 52, 27, 383376, tzinfo=dt.timezone.utc)
    local = dt.datetime(2026, 8, 3, 17, 52, 27, 383376)
    cols = [Column("CleanUrl", "Address", 400, kind="url"),
            Column("AccessedUtc", "Last visit (UTC)", 165, kind="time"),
            Column("AccessedLocal", "Last visit (LOCAL)", 165,
                   kind="time_local")]
    rows = [dict(_rows()[0], AccessedUtc=utc, AccessedLocal=local)]
    csv_path = str(tmp_path / "t.csv")
    F.write_view_csv(rows, cols, csv_path, {})
    header, line = _read(csv_path)
    xlsx_path = str(tmp_path / "t.xlsx")
    F.write_view_xlsx(rows, cols, xlsx_path, {"view_label": "History"})
    ws = openpyxl.load_workbook(xlsx_path).active
    for label, want in (("Last visit (UTC)", utc), ("Last visit (LOCAL)", local)):
        ci = header.index(label) + 1
        cell = ws.cell(row=2, column=ci)
        assert cell.is_date, (label, cell.value)
        assert ".000" in cell.number_format, (label, cell.number_format)
        serial = _stored_serial(xlsx_path, "%s2" % get_column_letter(ci))
        stored = dt.datetime(1899, 12, 30) + dt.timedelta(days=float(serial))
        naive = want.replace(tzinfo=None)
        assert abs(stored - naive) < dt.timedelta(microseconds=1), label
        from_csv = dt.datetime.fromisoformat(line[ci - 1]).replace(tzinfo=None)
        assert from_csv == naive, label


def test_the_view_workbook_is_streamed_not_held_in_memory(tmp_path):
    """I5: a normal openpyxl workbook holds every cell until save --
    ~27 MB per 1,000 rows at 85 columns, so a 100k-row view needed
    gigabytes. Streamed, memory stays flat as the rows grow. 1,000 rows x
    85 held in memory peak near 37 MB; the bound is well under that. The
    guards must still hold on the streamed path."""
    import tracemalloc
    heads = ["c%d" % i for i in range(85)]

    import datetime as dt
    stamp = dt.datetime(2026, 8, 3, 15, 52, 27, 383376,
                        tzinfo=dt.timezone.utc)

    def rows(n):
        for r in range(n):
            cells = [("https://example.test/%d/%d" % (r, i)) if i % 3
                     else RAW + r for i in range(85)]
            cells[1] = "=HYPERLINK(%d)" % r
            cells[2] = stamp
            yield cells

    path = str(tmp_path / "big.xlsx")
    tracemalloc.start()
    try:
        forensic_engine.write_simple_xlsx(path, "Big", heads, rows(1000),
                                          timestamp_columns=[2])
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert peak < 16 * 1024 * 1024, "peak %.1f MB" % (peak / 1e6)
    ws = openpyxl.load_workbook(path).active
    assert ws.max_row == 1001
    assert str(ws.cell(row=1001, column=1).value) == str(RAW + 999)
    formula = ws.cell(row=1001, column=2)
    assert formula.data_type == "s" and formula.value == "=HYPERLINK(999)"
    assert ws.cell(row=1001, column=3).is_date
    assert sum(1 for row in ws.iter_rows() for c in row
               if c.data_type == "f") == 0


def _case_maps():
    from webcache.engine.case_db import CaseDb
    key = CaseDb.record_keys(LIVE)[0]
    return ({key: ["synthetic-tag"]},
            {key: [{"comment": "synthetic note"}]})


def test_a_view_export_carries_the_cases_tags_and_comments(tmp_path):
    """I6: Tags and Analyst comments were always blank in a view file while
    the full export stated them -- a blank standing in for a known value."""
    tags, comments = _case_maps()
    path = str(tmp_path / "tagged.csv")
    out = _run(_job("view_csv", path, rows=_rows(), columns=COLS,
                    tags_map=tags, comments_map=comments,
                    scope={"view": "history", "view_label": "History",
                           "columns": [], "hidden_columns": [],
                           "rows_in_file": 2, "rows_in_view": 2,
                           "rows_in_record_set": 2, "is_subset": False,
                           "text_filter": "", "conditions": {}, "values": {},
                           "chip_text": "", "live_rows": 1, "carved_rows": 1,
                           "recovery_scan_run": True},
                    stamp={"exported_utc": "2026-09-26T00:00:00Z",
                           "evidence": {}}))
    assert out["failed"] == []
    header, live, carved = _read(path)
    assert live[header.index("Tags")] == "synthetic-tag"
    assert live[header.index("Analyst comments")] == "synthetic note"
    assert carved[header.index("Tags")] == ""


def test_export_view_hands_the_case_maps_to_the_job(qtbot, tmp_path,
                                                    monkeypatch):
    from webcache.shell.main_window import MainWindow
    tags, comments = _case_maps()
    w = MainWindow()
    qtbot.addWidget(w)
    view = w._ensure_view_built("history")
    view.set_records(_rows())
    jobs = []
    monkeypatch.setattr(w, "_tags_map", lambda: tags)
    monkeypatch.setattr(w, "_comments_map", lambda: comments)
    monkeypatch.setattr(w, "_choose_export_path",
                        lambda *a: str(tmp_path / "history.csv"))
    monkeypatch.setattr(w, "_start_export",
                        lambda job, card_key=None: jobs.append(job))
    w._export_view(view, "history")
    assert jobs[0].tags_map == tags
    assert jobs[0].comments_map == comments


def _sidecar(tmp_path, rows, cols, name):
    path = str(tmp_path / name)
    scope = {}
    F.write_view_csv(rows, cols, path, scope)
    return open(F.write_view_provenance(path, {}, scope, 0),
                encoding="utf-8-sig").read()


def test_the_sidecar_says_exactly_what_the_appended_block_is(tmp_path):
    """Review section 4: the old line claimed "every field these records
    carry" -- untrue of a fixed schema, and absurd beside a count of 0."""
    text = _sidecar(tmp_path, _rows(), COLS, "h.csv")
    canon = F.canonical_columns_for(_rows(), COLS)
    shown = len(forensic_engine.ENTRY_FIELDS) - len(canon)
    assert ("  Canonical fields : %d (the full ENTRY_FIELDS export schema, "
            "in schema order, minus the %d fields this view already shows; "
            "the same fixed column set every engine export writes. A blank "
            "cell means that record does not carry the field. Raw "
            "FILETIMEs, ids and carve provenance included.)"
            % (len(canon), shown)) in text
    agg = _sidecar(tmp_path, [{"CleanUrl": "example.test", "_kind": "domain",
                               "_record": None}], COLS, "d.csv")
    assert ("  Canonical fields : none (the rows in this view are "
            "aggregates, not single records, so no per-record fields "
            "apply)") in agg


def test_the_sidecar_names_a_non_entry_records_own_fields(tmp_path):
    rows, rec = _cookie_rows()
    text = _sidecar(tmp_path, rows, _cookie_columns(), "c.csv")
    assert "ENTRY_FIELDS export schema" not in text
    # Expires and LastModified are declared by the view AS those fields
    # (key == field), so the record's own copies are the ones not repeated.
    shown = {c.key for c in _cookie_columns() if c.field == c.key}
    assert shown == {"Expires", "LastModified"}
    assert "  Record fields    : %d" % (len(set(rec) - shown)) in text


def test_a_bool_column_reads_the_same_in_the_csv_and_the_xlsx(tmp_path):
    """A view's own bool column came out "True"/"False"
    in the CSV and "Yes"/blank in the xlsx of the same export -- and a
    blank there stood in for a KNOWN value. Both must write what the grid
    shows, "Yes"/"No", row for row."""
    cols = [Column("CleanUrl", "Address", 400, kind="url"),
            Column("InPrivate", "InPrivate", 80)]
    rows = [dict(_rows()[0], InPrivate=True), dict(_rows()[0], InPrivate=False)]
    csv_path = str(tmp_path / "b.csv")
    F.write_view_csv(rows, cols, csv_path, {})
    header, *body = _read(csv_path)
    xlsx_path = str(tmp_path / "b.xlsx")
    F.write_view_xlsx(rows, cols, xlsx_path, {"view_label": "History"})
    ws = openpyxl.load_workbook(xlsx_path).active
    ci = header.index("InPrivate")
    from_csv = [line[ci] for line in body]
    from_xlsx = [ws.cell(row=r + 2, column=ci + 1).value for r in range(len(rows))]
    assert from_csv == ["Yes", "No"], from_csv
    assert from_xlsx == from_csv, from_xlsx
